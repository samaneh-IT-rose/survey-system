# -*- coding: utf-8 -*-
"""خروجی اکسل رکوردهای نظرسنجی با openpyxl — برگه راست‌به‌چپ و ستون‌های پویا."""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import ORG_NAME
from ..constants import RATING_LABELS
from ..models import ContactOption, Question, Survey
from ..utils.jalali import gdate_to_jalali_str, today_jalali_str

HEADER_FILL = PatternFill("solid", fgColor="BA376D")
HEADER_FONT = Font(name="Vazirmatn", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Vazirmatn", size=10)
THIN = Side(style="thin", color="D9A9BF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)


def _fetch_filtered(db: Session, filters: dict) -> list[Survey]:
    from sqlalchemy import or_

    stmt = select(Survey).order_by(Survey.survey_date.desc(), Survey.id.desc())
    q = filters.get("q")
    if q:
        needle = f"%{q.strip()}%"
        stmt = stmt.where(
            or_(
                Survey.first_name.ilike(needle),
                Survey.last_name.ilike(needle),
                Survey.national_code.ilike(needle),
                Survey.phone.ilike(needle),
            )
        )
    if filters.get("province_id"):
        stmt = stmt.where(Survey.province_id == filters["province_id"])
    if filters.get("county_id"):
        stmt = stmt.where(Survey.county_id == filters["county_id"])
    if filters.get("contact_option_id"):
        stmt = stmt.where(Survey.contact_option_id == filters["contact_option_id"])
    from ..utils.jalali import jalali_to_gregorian, parse_jalali_str
    import datetime as dt

    pf = parse_jalali_str(filters.get("date_from") or "")
    if pf:
        gy, gm, gd = jalali_to_gregorian(*pf)
        stmt = stmt.where(Survey.survey_date >= dt.date(gy, gm, gd))
    pt = parse_jalali_str(filters.get("date_to") or "")
    if pt:
        gy, gm, gd = jalali_to_gregorian(*pt)
        stmt = stmt.where(Survey.survey_date <= dt.date(gy, gm, gd))
    return list(db.scalars(stmt))


def build_excel(db: Session, filters: dict) -> tuple[bytes, str]:
    questions = db.scalars(
        select(Question).where(Question.is_active.is_(True)).order_by(Question.display_order, Question.id)
    ).all()
    rows = _fetch_filtered(db, filters)

    wb = Workbook()
    ws = wb.active
    ws.title = "نظرسنجی"
    ws.sheet_view.rightToLeft = True

    headers = [
        "ردیف", "نام", "نام خانوادگی", "کد ملی", "شماره تماس",
        "استان", "شهرستان", "تاریخ نظرسنجی (شمسی)", "وضعیت تماس",
    ] + [f"سؤال {i + 1}: {q.text}" for i, q in enumerate(questions)] + [
        "زمان ثبت (شمسی)",
    ]

    # سطر عنوان گزارش
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=f"{ORG_NAME} — خروجی فرم‌های نظرسنجی ({today_jalali_str()})")
    title_cell.font = Font(name="Vazirmatn", size=12, bold=True, color="BA376D")
    title_cell.alignment = CENTER
    ws.row_dimensions[1].height = 26

    # سطر سرستون
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[2].height = 34

    # داده‌ها
    for i, s in enumerate(rows, start=1):
        answers = {a.question_id: a.value_label for a in s.answers}
        contact_text = s.contact_option.text if s.contact_option else "—"
        values = [
            i, s.first_name, s.last_name, s.national_code, s.phone,
            s.province.name, s.county.name, gdate_to_jalali_str(s.survey_date), contact_text,
        ] + [answers.get(q.id, "") for q in questions] + [
            gdate_to_jalali_str(s.created_at.date()) if s.created_at else "",
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=i + 2, column=col, value=v)
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = CENTER if col in (1, 4, 5, 8) else RIGHT

    # عرض ستون‌ها
    base_widths = [6, 14, 16, 14, 14, 12, 14, 16, 24]
    for col, w in enumerate(base_widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    for i in range(len(questions)):
        ws.column_dimensions[get_column_letter(len(base_widths) + i + 1)].width = 16
    ws.column_dimensions[get_column_letter(len(headers))].width = 16
    ws.freeze_panes = "C3"

    bio = io.BytesIO()
    wb.save(bio)
    filename = f"survey-export-{today_jalali_str().replace('/', '-')}.xlsx"
    return bio.getvalue(), filename
